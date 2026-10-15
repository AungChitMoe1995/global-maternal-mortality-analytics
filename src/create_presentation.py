"""Build a 15-slide briefing with editable tables and native data charts."""
from pathlib import Path
import math
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from research_io import load_results
from presentation_cache import sync_numeric_caches

NAVY='18334A'; TEAL='087F78'; MUTED='526579'; RED='BD5829'
INCOMES=['Low income','Lower middle income','Upper middle income','High income']
COLORS=['BD5829','7C63A8','168B80','315D89']


def create_deck(base_dir):
    root=Path(base_dir)
    r,audit=load_results(root)
    df=pd.read_csv(root/'data/processed/maternal_mortality_analytical_master.csv')
    y=df.loc[df.year.eq(2023)&df.is_balanced_sample.eq(1)]
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def text(slide,copy,x,y,w,h,size=20,color=NAVY,bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=box.text_frame;tf.word_wrap=True
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for i,line in enumerate(copy.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=Pt(12)
        return box
    def slide(title):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
        text(s,title,.7,.42,11.95,.9,30,bold=True)
        text(s,f'Maternal mortality research · WHO/UN MMEIG 2025 and World Bank WDI · {len(prs.slides)}/15',.7,7.02,11.95,.25,10,MUTED)
        s.notes_slide.notes_text_frame.text=('Source: reports/results.json; analytical CSV SHA-256 '+r['data_sha256']+'\nWHO report: https://www.who.int/publications/i/item/9789240108462\nWorld Bank: https://databank.worldbank.org/source/world-development-indicators\nAll regression estimates are exploratory associations with modeled ecological outcomes.')
        return s
    def body(s,lines,top=1.6,size=22):
        text(s,'\n'.join(lines),.8,top,11.7,5.1,size)
    def caption(s,copy): text(s,copy,.8,6.15,11.75,.7,15,MUTED)
    def table(s,heads,rows,x=.8,top=1.6,w=11.7,h=3.8,size=17,widths=None):
        shape=s.shapes.add_table(len(rows)+1,len(heads),Inches(x),Inches(top),Inches(w),Inches(h))
        if widths:
            for column,width in zip(shape.table.columns,widths): column.width=Inches(width)
        for i,row in enumerate([heads]+rows):
            for j,value in enumerate(row):
                cell=shape.table.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.12);cell.margin_right=Inches(.1)
                cell.text_frame.word_wrap=True
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(NAVY if i==0 else ('F0F5F7' if i%2 else 'FFFFFF'))
                for p in cell.text_frame.paragraphs:
                    p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=i==0
                    p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else NAVY)
        return shape
    def chart(s,data,kind,x=.8,top=1.55,w=11.65,h=4.45,legend=True,x_range=None):
        c=s.shapes.add_chart(kind,Inches(x),Inches(top),Inches(w),Inches(h),data).chart
        c.has_title=False
        c.has_legend=legend
        if legend:
            c.legend.position=XL_LEGEND_POSITION.BOTTOM;c.legend.include_in_layout=False;c.legend.font.size=Pt(12)
        for axis in [c.category_axis,c.value_axis]:
            axis.tick_labels.font.name='Aptos';axis.tick_labels.font.size=Pt(12)
        c.value_axis.major_gridlines.format.line.color.rgb=RGBColor.from_string('DBE4EA')
        if kind==XL_CHART_TYPE.XY_SCATTER:
            c.value_axis.minimum_scale=math.floor(math.log10(y.mmr_pt.min()))-.1
            c.value_axis.maximum_scale=math.ceil(math.log10(y.mmr_pt.max()))+.1
            if x_range:
                c.category_axis.minimum_scale=x_range[0];c.category_axis.maximum_scale=x_range[1]
        for i,series in enumerate(c.series):
            color=RGBColor.from_string(COLORS[i%4]);series.format.line.color.rgb=color
            if kind==XL_CHART_TYPE.XY_SCATTER:
                series.format.line.fill.background()
                series.marker.style=XL_MARKER_STYLE.CIRCLE
                series.marker.format.fill.solid();series.marker.format.fill.fore_color.rgb=color
                series.marker.format.line.color.rgb=color;series.marker.size=4
            elif kind==XL_CHART_TYPE.COLUMN_CLUSTERED:
                series.format.fill.solid();series.format.fill.fore_color.rgb=color
        return c

    s=slide('Global maternal mortality, 2000–2023')
    text(s,'Trends, inequalities and\nexploratory associations',.8,1.8,11.7,1.6,38,bold=True)
    text(s,f"{r['sample']['all_countries']} countries and territories\n{r['sample']['balanced_countries']} countries in the complete balanced model",.8,3.8,11.7,1.2,24,TEAL)
    text(s,'Dr. Aung Chit Moe, MBBS\nHealthcare data analytics portfolio',.8,5.45,11.7,.8,19,MUTED)
    s=slide('Research question and analytical scope')
    body(s,['How did estimated maternal mortality change across countries?',
            'How do country-level inequalities differ by fixed income group?',
            'How are estimated outcomes associated with economic resources and national fertility indicators?',
            'Unit of analysis: country and calendar year. The data do not identify individual clinical risk.'])
    s=slide('The outcome already depends on economic inputs')
    body(s,['WHO/UN MMEIG combines mortality observations with statistical estimation.',
            'BMat uses GDP, general fertility and skilled birth attendance. These inputs become more influential when direct observations are sparse.',
            'A GDP regression on MMEIG estimates is partly dependent on how the outcome was constructed.',
            'The portfolio presents descriptive trends and exploratory associations. It does not independently establish economic or fertility effects.'],size=21)
    caption(s,'WHO 2025 report, sections 3.2.4 and 3.4.2. Outcome intervals are marginal 80% uncertainty limits.')
    s=slide('Data coverage and integrity')
    sample=r['sample']
    table(s,['Sample','Countries','Country-years'],[
        ['Full outcome grid',sample['all_countries'],f"{sample['all_rows']:,}"],
        ['Balanced complete panel',sample['balanced_countries'],f"{sample['balanced_rows']:,}"],
        ['All complete country-years',sample['complete_countries'],f"{sample['complete_rows']:,}"],
        ['Calculated integrity audit',audit['status'],str(len(audit['checks']))+' checks']],h=3.4,size=20)
    caption(s,'The audit checks keys, bounds, logarithms, sample flags, storage consistency and cached hashes. Passing does not establish scientific validity.')
    s=slide('Country averages and aggregate mortality differ')
    data=CategoryChartData();data.categories=[x['year'] for x in r['global_trends']]
    for key,label in [('median_mmr','Country median'),('aggregate_mmr','Derived aggregate MMR')]:
        data.add_series(label,[x[key] for x in r['global_trends']])
    chart(s,data,XL_CHART_TYPE.LINE)
    caption(s,'MMR: deaths per 100,000 live births. The derived ratio uses summed country median death estimates and births, so it may differ from the official global posterior result.')
    s=slide('Income-group differences persist')
    data=CategoryChartData();data.categories=list(range(2000,2024))
    for income in INCOMES:
        data.add_series(income,[x['median_mmr'] for x in r['income_trends'] if x['income']==income])
    chart(s,data,XL_CHART_TYPE.LINE)
    caption(s,'Country-median MMR using all available countries. Income labels are fixed at acquisition, not historical classifications. Cook Islands is unclassified and omitted from these four group lines.')
    s=slide('GDP and estimated mortality, 2023')
    data=XyChartData()
    for income in INCOMES:
        series=data.add_series(income)
        for _,row in y.loc[y.income_level_name.eq(income)].iterrows(): series.add_data_point(math.log10(row.gdp_pcap_ppp),math.log10(row.mmr_pt))
    chart(s,data,XL_CHART_TYPE.XY_SCATTER,x_range=(math.floor(math.log10(y.gdp_pcap_ppp.min())),math.ceil(math.log10(y.gdp_pcap_ppp.max()))))
    caption(s,'Horizontal: log10 GDP per capita, constant 2021 PPP dollars. Vertical: log10 MMR. Each point is a balanced-sample country. GDP also informs MMEIG estimates.')
    s=slide('National fertility indicators and mortality')
    for j,(var,label) in enumerate([('tfr','Total fertility: births/woman'),('ado_fert_rate','Adolescent fertility: births/1,000')]):
        data=XyChartData()
        for income in INCOMES:
            series=data.add_series(income)
            for _,row in y.loc[y.income_level_name.eq(income)].iterrows(): series.add_data_point(row[var],math.log10(row.mmr_pt))
        chart(s,data,XL_CHART_TYPE.XY_SCATTER,x=.8+j*6,w=5.7,h=3.9,top=1.65,legend=False,x_range=(0,math.ceil(y[var].max())+.5))
        text(s,label,.9+j*6,5.6,5.5,.45,17)
    caption(s,'Vertical axes: log10 MMR, 2023 balanced sample. National fertility does not measure individual parity, birth spacing or pregnancy risk. Demographic inputs also inform upstream estimation.')
    s=slide('Health expenditure measures several financing sources')
    data=XyChartData()
    for income in INCOMES:
        series=data.add_series(income)
        for _,row in y.loc[y.income_level_name.eq(income)].iterrows(): series.add_data_point(row.chex_gdp_pct,math.log10(row.mmr_pt))
    chart(s,data,XL_CHART_TYPE.XY_SCATTER,x=.8,w=6.3,h=4.2,x_range=(0,math.ceil(y.chex_gdp_pct.max())+1))
    data=CategoryChartData();data.categories=[x['quartile'] for x in r['spending_quartiles']]
    data.add_series('Median MMR',[x['median_mmr'] for x in r['spending_quartiles']])
    chart(s,data,XL_CHART_TYPE.COLUMN_CLUSTERED,x=7.4,w=5.05,h=4.2,legend=False)
    caption(s,'Left: expenditure % GDP versus log10 MMR. Right: calculated expenditure quartiles versus country-median MMR. CHE includes public, private and external financing, rather than government priority alone.')
    s=slide('Balanced-panel exploratory coefficients')
    rows=[]
    for c in r['models']['balanced']['coefficients']:
        label={'log_gdp_pcap_ppp':'Log GDP PPP','chex_gdp_pct':'Health exp. % GDP','tfr':'Total fertility','ado_fert_rate':'Adolescent fertility','urban_pct':'Urban population %'}[c['term']]
        rows.append([label,f"{c['coefficient']:.4f}",f"{c['ci_low']:.4f} to {c['ci_high']:.4f}"])
    table(s,['Predictor','Coefficient','95% confidence interval'],rows,size=18)
    caption(s,'Outcome: log MMR. Country and year fixed effects, equal country-year weighting and country-clustered t inference. CIs are conditional on the model and do not propagate all MMEIG uncertainty.')
    s=slide('Sensitivity to sample and specification')
    rows=[]
    for key in ['balanced','unbalanced','pre_pandemic','lower_bound','upper_bound','first_difference']:
        m=r['models'][key];c=m['coefficients'][0]
        rows.append([m['label'],f"{m['n_obs']:,}",f"{c['coefficient']:.3f}",f"{c['ci_low']:.3f} to {c['ci_high']:.3f}"])
    table(s,['GDP slope specification','Rows','Coefficient','95% CI'],rows,size=16,h=4.25,widths=[4.65,1.2,1.8,4.05])
    caption(s,'Marginal outcome-bound replacements are sensitivity fits, not joint posterior draws. Full tables also report four income-subset models and all other predictors.')
    s=slide('Diagnostics identify concerns without removing bias')
    rows=[[x['label'],f"{x['vif']:.2f}"] for x in r['vif']]
    table(s,['Within-panel predictor','VIF'],rows,w=7.1,h=3.9,size=17,widths=[5.5,1.6])
    d=r['diagnostics']
    text(s,f"Residual lag-1 correlation\n{d['residual_lag1_correlation']:.3f}\n\nGDP slope after one-country omission\n{d['leave_one_out_gdp_min']:.3f} to {d['leave_one_out_gdp_max']:.3f}",8.3,1.8,4,3.7,20)
    caption(s,'Clustered standard errors address within-country dependence under cluster assumptions. Upstream shared modeling, changing confounding and influential countries remain relevant.')
    s=slide('The balanced sample excludes different settings')
    rows=[[x['year'],x['group'],x['n_countries'],f"{x['median_mmr']:.1f}"] for x in r['selection']]
    table(s,['Year','Sample group','Countries','Median MMR'],rows,h=3.4,size=20)
    caption(s,'Excluded settings include Afghanistan, Somalia, South Sudan, Syria and Yemen. Complete-case analysis does not eliminate selection bias. Full descriptive outputs retain all available countries.')
    s=slide('Interpretation and remaining research needs')
    body(s,['The study describes national trends and inequalities and examines exploratory conditional associations.',
            'It cannot identify clinical mechanisms, individual risk, or the effects of changing expenditure or fertility.',
            'A source-quality-restricted analysis needs defensible country-period CRVS classifications.',
            'Full uncertainty propagation requires suitable joint posterior draws. These tasks remain open.'],size=22)
    s=slide('Reproducible portfolio outputs')
    body(s,['One audited input snapshot feeds calculated tables, figures, the notebook, dashboard and this briefing.',
            'Run: python src/run_pipeline.py',
            'Check: python -m pytest -q',
            'Source code: github.com/AungChitMoe1995/global-maternal-mortality-analytics',
            'AI tools assisted implementation. I directed the research questions and reviewed the outputs within my current technical understanding.'],size=21)
    path=root/'docs/presentation/maternal_mortality_research_briefing.pptx';path.parent.mkdir(exist_ok=True)
    prs.save(path)
    sync_numeric_caches(path)
    print('Built',len(prs.slides),'slides with editable charts and tables')


if __name__=='__main__':
    create_deck(Path(__file__).resolve().parents[1])
